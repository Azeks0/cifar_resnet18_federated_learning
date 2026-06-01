from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
RESULTS_DIR = PROJECT_ROOT / "outputs" / "results"
PLOTS_DIR = PROJECT_ROOT / "outputs" / "plots"
PAPER_DIR = PROJECT_ROOT / "paper"
REPORT_PATH = PAPER_DIR / "resnet18_cifar10_baseline_report.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
GRAY = RGBColor(89, 89, 89)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
BORDER = "D9DEE7"


def set_run_font(run, *, size=None, color=None, bold=None, italic=None, name="Calibri") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(23)
    title.font.color.rgb = INK
    title.font.bold = True
    title.paragraph_format.space_after = Pt(4)

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1


def paragraph_bottom_border(paragraph, color="2E74B5", size="8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "bottom": bottom, "start": start, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER) -> None:
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table._tbl.tblPr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            row.cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(row.cells[idx])


def format_cell(cell, value, *, bold=False, size=9.4, align=None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.1
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(str(value))
    set_run_font(run, size=size, color=INK, bold=bold)


def add_table(doc: Document, dataframe: pd.DataFrame, columns: list[tuple[str, str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    set_table_borders(table)
    set_table_geometry(table, widths)
    for idx, (_, label) in enumerate(columns):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_FILL)
        format_cell(cell, label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for _, row in dataframe.iterrows():
        cells = table.add_row().cells
        for idx, width in enumerate(widths):
            cells[idx].width = Inches(width)
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[idx])
        for idx, (column, _) in enumerate(columns):
            format_cell(cells[idx], row[column])
    doc.add_paragraph()


def add_paragraph(doc: Document, text: str = ""):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_callout(doc: Document, label: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    set_table_borders(table, color="E2E8F0")
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, size=10.5, color=INK, bold=True)
    body_run = paragraph.add_run(text)
    set_run_font(body_run, size=10.5, color=INK)
    doc.add_paragraph()


def add_code_block(doc: Document, code: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    set_table_borders(table, color="E2E8F0")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "FAFBFC")
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    for line_no, line in enumerate(code.splitlines()):
        if line_no:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(run, name="Courier New", size=8.8, color=INK)
    doc.add_paragraph()


def add_picture(doc: Document, path: Path, caption: str) -> None:
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(6.1))
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption_paragraph.add_run(caption)
    set_run_font(caption_run, size=9, color=GRAY, italic=True)


def maybe_results() -> tuple[pd.DataFrame | None, pd.DataFrame | None]:
    summary_path = RESULTS_DIR / "summary.csv"
    history_path = RESULTS_DIR / "history.csv"
    summary = pd.read_csv(summary_path) if summary_path.exists() else None
    history = pd.read_csv(history_path) if history_path.exists() else None
    return summary, history


def build_report() -> None:
    PAPER_DIR.mkdir(parents=True, exist_ok=True)
    summary, history = maybe_results()

    doc = Document()
    configure_document(doc)

    doc.add_paragraph()
    title = doc.add_paragraph(style="Title")
    title.add_run("Centralized Baseline for Federated Learning")
    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run("ResNet-18 on CIFAR-10")
    set_run_font(subtitle_run, size=13, color=GRAY)

    metadata = [
        ("Project", "Federated Learning assignment - baseline step"),
        ("Dataset", "CIFAR-10"),
        ("Architecture", "ResNet-18 adapted for 32x32 images"),
        ("Generated", date.today().strftime("%B %d, %Y")),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        lr = p.add_run(f"{label}: ")
        set_run_font(lr, bold=True, color=INK)
        vr = p.add_run(value)
        set_run_font(vr, color=INK)
    rule = doc.add_paragraph()
    paragraph_bottom_border(rule)

    doc.add_heading("Abstract", level=1)
    add_paragraph(
        doc,
        "This report documents the centralized baseline required before implementing federated-learning strategies. "
        "The task is image classification on CIFAR-10 using a full-precision ResNet-18. The resulting model, training "
        "curves, validation accuracy, and test accuracy form the reference point for later FedAvg or custom FL strategies."
    )
    add_callout(
        doc,
        "Baseline role",
        "This is not the federated method itself. It is the standard centralized experiment that establishes the accuracy "
        "and loss target for later federated runs on the same dataset and architecture.",
    )

    doc.add_heading("1. Problem", level=1)
    add_paragraph(
        doc,
        "The federated-learning project asks for a classification baseline before client-server training is introduced. "
        "A centralized baseline is useful because all training samples are available to a single optimizer, so it gives a "
        "clean upper reference for convergence and generalization under the selected model architecture."
    )

    doc.add_heading("2. Dataset", level=1)
    add_paragraph(
        doc,
        "CIFAR-10 contains 60,000 RGB images of size 32 by 32 pixels across ten classes: airplane, automobile, bird, cat, "
        "deer, dog, frog, horse, ship, and truck. The project uses 45,000 training samples, 5,000 validation samples, and "
        "the official 10,000-image test split. Training augmentation consists of random crop with padding 4 and random "
        "horizontal flip. Evaluation uses only normalization."
    )

    doc.add_heading("3. Algorithm and Architecture", level=1)
    add_paragraph(
        doc,
        "The classifier is a ResNet-18 adapted for CIFAR resolution. The first layer is a 3x3 convolution with stride 1, and "
        "the ImageNet-style initial max-pooling layer is removed so early spatial information is not discarded on 32x32 inputs. "
        "The residual stages use two BasicBlock modules each, followed by global average pooling and a ten-way linear classifier."
    )
    add_code_block(
        doc,
        "for epoch = 1 ... E:\n"
        "    train ResNet-18 on augmented CIFAR-10 batches\n"
        "    evaluate validation loss and accuracy\n"
        "    save checkpoint when validation accuracy improves\n"
        "load best validation checkpoint\n"
        "evaluate train, validation, and test accuracy/loss\n"
        "export curves, confusion matrix, and summary table"
    )

    doc.add_heading("4. Experimental Setup", level=1)
    setup = pd.DataFrame(
        [
            ["Model", "CIFAR-style ResNet-18"],
            ["Optimizer", "SGD with momentum 0.9 and Nesterov acceleration"],
            ["Learning rate", "0.1 with cosine annealing"],
            ["Weight decay", "0.0005"],
            ["Batch size", "128"],
            ["Epochs", "30 by default"],
            ["Metric focus", "Accuracy and cross-entropy loss"],
        ],
        columns=["Setting", "Value"],
    )
    add_table(doc, setup, [("Setting", "Setting"), ("Value", "Value")], [2.2, 4.3])

    doc.add_heading("5. Results", level=1)
    if summary is not None:
        is_smoke = bool(
            int(summary.iloc[0].get("epochs", 0)) < 30
            or int(summary.iloc[0].get("train_samples", 0)) < 45000
        )
        if is_smoke:
            add_callout(
                doc,
                "Result scope",
                "The table below reflects the latest available smoke-test run, not the full 30-epoch baseline. "
                "Run the default command and rebuild the paper to replace these values with full baseline results.",
            )
        display = summary.copy()
        for column in ["train_accuracy", "val_accuracy", "test_accuracy"]:
            if column in display.columns:
                display[column] = display[column].map(lambda value: f"{value * 100:.2f}%")
        for column in ["train_loss", "val_loss", "test_loss"]:
            if column in display.columns:
                display[column] = display[column].map(lambda value: f"{value:.3f}")
        result_columns = [
            ("epochs", "Epochs"),
            ("train_samples", "Train"),
            ("val_samples", "Val"),
            ("test_samples", "Test"),
            ("val_accuracy", "Val Acc."),
            ("test_accuracy", "Test Acc."),
            ("test_loss", "Test Loss"),
        ]
        add_table(doc, display, result_columns, [0.7, 0.8, 0.7, 0.7, 0.9, 0.9, 0.9])
        if history is not None:
            add_picture(doc, PLOTS_DIR / "accuracy_curve.png", "Figure 1. Training and validation accuracy by epoch.")
            add_picture(doc, PLOTS_DIR / "loss_curve.png", "Figure 2. Training and validation loss by epoch.")
        add_picture(doc, PLOTS_DIR / "confusion_matrix.png", "Figure 3. Test confusion matrix.")
    else:
        add_paragraph(
            doc,
            "No local training summary was present when this document was generated. Running `scripts/run_baseline.py` will "
            "populate the results table, learning curves, and confusion matrix automatically; rebuilding this document then "
            "embeds those artifacts."
        )

    doc.add_heading("6. Discussion", level=1)
    add_paragraph(
        doc,
        "For the federated-learning project, the important property of this baseline is repeatability. The same CIFAR-10 split, "
        "model, optimizer, and metrics can be reused when later replacing centralized data access with client partitions. Any "
        "federated method should be compared against the centralized test accuracy and loss reported here."
    )
    add_paragraph(
        doc,
        "A centralized ResNet-18 should converge faster and usually reach a higher final accuracy than non-IID federated runs "
        "because every batch is sampled from the full training distribution. If a future FL strategy approaches this baseline "
        "with fewer communication rounds or lower instability, it is evidence that the strategy handles decentralized data well."
    )

    doc.add_heading("7. Reproducibility", level=1)
    add_paragraph(doc, "Run the full baseline with:")
    add_code_block(doc, "PYTHONPATH=src python3 scripts/run_baseline.py --config configs/default.json")
    add_paragraph(doc, "Run a quick smoke test with:")
    add_code_block(
        doc,
        "PYTHONPATH=src python3 scripts/run_baseline.py --config configs/default.json --epochs 1 "
        "--max-train-samples 512 --max-val-samples 256 --max-test-samples 256",
    )

    doc.add_heading("References", level=1)
    for reference in [
        "[1] Federated Learning Project Handout, Revision 1.",
        "[2] K. He, X. Zhang, S. Ren, and J. Sun. Deep residual learning for image recognition. CVPR, 2016.",
        "[3] A. Krizhevsky. Learning multiple layers of features from tiny images. Technical report, 2009.",
        "[4] B. McMahan, E. Moore, D. Ramage, S. Hampson, and B. Aguera y Arcas. Communication-efficient learning of deep networks from decentralized data. AISTATS, 2017.",
    ]:
        add_paragraph(doc, reference)

    doc.save(REPORT_PATH)
    print(REPORT_PATH)


if __name__ == "__main__":
    build_report()
